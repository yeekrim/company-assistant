"""업로드/삭제 API가 표 저장소와 연결되어 있는지 확인. (인증과 Chroma는 가짜로 대체)"""
import io

import docx
import pytest
from fastapi.testclient import TestClient

from app.api import documents
from app.core.config import settings
from app.main import app
from app.services import document_service
from app.tables import store


class FakeUser:
    id = 1
    company_id = 7
    role = "admin"


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "TABLE_STORE_DIR", str(tmp_path))
    monkeypatch.setattr(document_service, "add_documents", lambda *a, **k: None)
    monkeypatch.setattr(documents, "delete_document", lambda *a, **k: None)
    app.dependency_overrides[documents.require_admin] = lambda: FakeUser()
    yield TestClient(app)
    app.dependency_overrides.clear()


def docx_with_table() -> bytes:
    d = docx.Document()
    t = d.add_table(rows=3, cols=2)
    for i, row in enumerate([["부서", "인원"], ["영업팀", "12"], ["개발팀", "20"]]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def upload(client, name, data):
    return client.post("/api/documents/upload", files=[("files", (name, data))]).json()


def test_upload_saves_tables(client):
    res = upload(client, "인원.docx", docx_with_table())
    assert res["uploaded"][0]["tables"] == 1
    [t] = store.list_tables(FakeUser.company_id)
    assert t["doc_name"] == "인원.docx"
    assert t["row_count"] == 2


def test_upload_without_tables(client):
    res = upload(client, "a.txt", "표 없음".encode())
    assert res["uploaded"][0]["tables"] == 0
    assert res["errors"] == []


def test_table_failure_does_not_fail_upload(client, monkeypatch):
    def boom(*a, **k):
        raise RuntimeError("추출 실패")
    monkeypatch.setattr(document_service, "extract_tables", boom)
    res = upload(client, "인원.docx", docx_with_table())
    assert res["errors"] == []
    assert res["uploaded"][0]["tables"] == 0


def test_delete_removes_tables(client):
    upload(client, "인원.docx", docx_with_table())
    r = client.delete("/api/documents/인원.docx")
    assert r.status_code == 200
    assert store.list_tables(FakeUser.company_id) == []
