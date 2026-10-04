from datetime import date

import docx

from app.tables.extractor import assemble, extract_tables


# ── assemble 단위 테스트 ──────────────────────────────────────────

def test_numbers_with_commas_become_int():
    [t] = assemble([("", [["이름", "금액"], ["김", "12,300"], ["이", "1,000,000"]])])
    assert [r[1] for r in t.rows] == [12300, 1000000]


def test_decimals_become_float():
    [t] = assemble([("", [["이름", "연차"], ["김", "9.5"], ["이", "3"]])])
    assert [r[1] for r in t.rows] == [9.5, 3.0]


def test_dates_in_various_formats():
    [t] = assemble([("", [["일자"], ["2026-01-05"], ["2026.02.10"], ["2026/3/1"]])])
    assert [r[0] for r in t.rows] == [date(2026, 1, 5), date(2026, 2, 10), date(2026, 3, 1)]


def test_mixed_column_stays_text():
    [t] = assemble([("", [["코드"], ["123"], ["A12"]])])
    assert [r[0] for r in t.rows] == ["123", "A12"]


def test_invalid_date_falls_back_to_text():
    [t] = assemble([("", [["일자"], ["2026-02-30"], ["2026-03-01"]])])
    assert [r[0] for r in t.rows] == ["2026-02-30", "2026-03-01"]


def test_empty_cells_become_none():
    [t] = assemble([("", [["이름", "금액"], ["김", ""], ["이", "500"]])])
    assert [r[1] for r in t.rows] == [None, 500]


def test_continued_table_across_pages_is_merged():
    header = ["이름", "금액"]
    raw = [("부록 1", [header, ["김", "100"]]), ("", [header, ["이", "200"]])]
    [t] = assemble(raw)
    assert t.title == "부록 1"
    assert len(t.rows) == 2


def test_different_headers_are_separate_tables():
    raw = [("", [["이름", "금액"], ["김", "100"]]), ("", [["부서", "인원"], ["영업", "3"]])]
    assert len(assemble(raw)) == 2


def test_header_only_table_is_skipped():
    assert assemble([("", [["이름", "금액"]])]) == []


def test_blank_and_duplicate_headers_are_named():
    [t] = assemble([("", [["금액", "", "금액"], ["1", "2", "3"]])])
    assert t.columns == ["금액", "열2", "금액_2"]


def test_short_rows_are_padded():
    [t] = assemble([("", [["a", "b", "c"], ["x", "1"]])])
    assert t.rows == [["x", 1, None]]


# ── extract_tables (파일) ──────────────────────────────────────────

def test_extract_docx_tables(tmp_path):
    d = docx.Document()
    d.add_paragraph("본문")
    tbl = d.add_table(rows=3, cols=2)
    for i, row in enumerate([["부서", "인원"], ["영업팀", "12"], ["개발팀", "20"]]):
        for j, v in enumerate(row):
            tbl.cell(i, j).text = v
    path = tmp_path / "doc.docx"
    d.save(path)

    [t] = extract_tables(str(path))
    assert t.columns == ["부서", "인원"]
    assert t.rows == [["영업팀", 12], ["개발팀", 20]]


def test_txt_has_no_tables(tmp_path):
    path = tmp_path / "a.txt"
    path.write_text("이름 | 금액", encoding="utf-8")
    assert extract_tables(str(path)) == []
